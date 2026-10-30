import importlib.util, io, pathlib, tempfile, threading, unittest, http.client, shutil, sys
from docx import Document
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from pypdf import PdfReader
import re, unicodedata, subprocess

spec=importlib.util.spec_from_file_location('office_server',pathlib.Path(__file__).parents[1]/'office-service/server.py');module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module)
class OfficeServiceTest(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=pathlib.Path(self.tmp.name);self.token='test-only-'+'a'*40
  self.server=module.OfficeServer(('127.0.0.1',0),module.Config(self.token,self.root/'jobs',shutil.which('soffice'),timeout=40));self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
 def tearDown(self):
  self.server.shutdown();self.server.server_close();self.tmp.cleanup()
 def request(self,path,body=b'',auth=True,headers=None,method='POST'):
  client=http.client.HTTPConnection(*self.server.server_address,timeout=60);h={'Content-Type':'application/octet-stream'}
  if auth:h['Authorization']='Bearer '+self.token
  h.update(headers or {});client.request(method,path,body=body,headers=h);response=client.getresponse();data=response.read();code=response.status;client.close();return code,data
 def clean(self):
  self.assertTrue(self.server.slot.acquire(timeout=3))
  try:self.assertEqual(list((self.root/'jobs').iterdir()),[])
  finally:self.server.slot.release()
 def test_auth_limits_bad_file_cleanup(self):
  self.assertEqual(self.request('/convert?ext=docx',b'x',auth=False)[0],401)
  self.assertEqual(self.request('/convert?ext=exe',b'x')[0],415)
  self.assertEqual(self.request('/convert?ext=docx',b'x',headers={'Content-Length':str(501*1024*1024)})[0],413)
  self.assertEqual(self.request('/convert?ext=docx',b'not an Office document')[0],422)
  self.clean()
  self.assertEqual(self.request('/health',method='GET')[0],200)
 def test_real_docx_formula_table_image_and_two_slide_pptx(self):
  image=self.root/'picture.png';Image.new('RGB',(80,40),(20,120,200)).save(image)
  doc=Document();doc.add_heading('WORD PREVIEW TEST',0);doc.add_paragraph('Before formula')
  equation=OxmlElement('m:oMath');run=OxmlElement('m:r');text=OxmlElement('m:t');text.text='x² + y² = 1';run.append(text);equation.append(run);doc.add_paragraph()._p.append(equation)
  table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='TABLE CELL';table.cell(1,1).text='ROW TWO';doc.add_picture(str(image));doc.add_page_break();doc.add_paragraph('FINAL PAGE MARKER');path=self.root/'input.docx';doc.save(path)
  code,data=self.request('/convert?ext=docx',path.read_bytes());self.assertEqual(code,200,data[:200]);reader=PdfReader(io.BytesIO(data));self.assertEqual(len(reader.pages),2);text=''.join(p.extract_text() for p in reader.pages);self.assertIn('TABLE CELL',text);self.assertIn('FINAL PAGE MARKER',text);self.assertIn('x2+y2=1',re.sub(r'\s+','',unicodedata.normalize('NFKC',text)));self.assertTrue(any(list(p.images) for p in reader.pages))
  ppt=Presentation()
  for title in ['FIRST SLIDE','SECOND SLIDE']:
   slide=ppt.slides.add_slide(ppt.slide_layouts[6]);slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text=title;slide.shapes.add_picture(str(image),Inches(1), Inches(2))
  path=self.root/'input.pptx';ppt.save(path);code,data=self.request('/convert?ext=pptx',path.read_bytes());self.assertEqual(code,200,data[:200]);reader=PdfReader(io.BytesIO(data));self.assertEqual(len(reader.pages),2);self.assertIn('SECOND SLIDE',reader.pages[1].extract_text());self.assertTrue(list(reader.pages[0].images));self.clean()
  legacy=self.root/'legacy';legacy.mkdir()
  for ext,filter_name in [('doc','MS Word 97'),('ppt','MS PowerPoint 97')]:
   original=self.root/('input.docx' if ext=='doc' else 'input.pptx')
   subprocess.run([shutil.which('soffice'),'-env:UserInstallation='+(self.root/'legacy-profile').as_uri(),'--headless','--convert-to',ext+':'+filter_name,'--outdir',str(legacy),str(original)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=40)
   code,data=self.request('/convert?ext='+ext,(legacy/('input.'+ext)).read_bytes());self.assertEqual(code,200,data[:200]);self.assertEqual(len(PdfReader(io.BytesIO(data)).pages),2);self.clean()
  code,data=self.request('/convert?ext=doc',br'{\rtf1\ansi RTF SOURCE}');self.assertEqual(code,200,data[:200]);self.assertIn('RTF SOURCE',PdfReader(io.BytesIO(data)).pages[0].extract_text());self.clean()
if __name__=='__main__':unittest.main()
